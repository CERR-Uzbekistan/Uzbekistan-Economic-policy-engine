"""Materialize deterministic test-only XLSX cell projections, never source workbooks."""

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from models.readiness import semantic_hash
from openpyxl import Workbook, load_workbook


def cell_view(fixture):
    result = {}
    for title, rows in fixture["sheets"].items():
        rows = [
            [float(format(value, ".16g")) if isinstance(value, float) else value for value in row]
            for row in rows
        ]
        while rows and all(value is None for value in rows[-1]):
            rows.pop()
        result[title] = rows
    return result


WORKBOOKS = ["cbu-bulletin-cells-v1.json", "world-bank-commodity-cells-v1.json"]


def prepare_workbooks(root, generated, read):
    output = {}
    for name in WORKBOOKS:
        fixture = read(root / "mcp_server/tests/fixtures" / name)
        if fixture["schema_version"] != "research-workbook-test-cells-v1":
            raise ValueError("research_workbook_fixture_schema_incompatible")
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.created = datetime(2000, 1, 1)  # noqa: DTZ001 -- XLSX properties require a naive fixed timestamp
        for title, rows in cell_view(fixture).items():
            sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(row)
        memory = io.BytesIO()
        workbook.save(memory)
        destination = generated / Path(fixture["source_path"]).name
        with (
            zipfile.ZipFile(memory) as source,
            zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED) as target,
        ):
            for entry in sorted(source.namelist()):
                body = source.read(entry)
                if entry == "docProps/core.xml":
                    body = re.sub(
                        rb"(<dcterms:modified[^>]*>)[^<]+", rb"\g<1>2000-01-01T00:00:00Z", body
                    )
                info = zipfile.ZipInfo(entry, (2000, 1, 1, 0, 0, 0))
                info.create_system = 3
                info.external_attr = 0o600 << 16
                target.writestr(info, body)
        output[destination.name] = semantic_hash(cell_view(fixture))
    return output


def validate_workbooks(root, generated, read):
    output = {}
    for name in WORKBOOKS:
        fixture = read(root / "mcp_server/tests/fixtures" / name)
        path = generated / Path(fixture["source_path"]).name
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = {
            title: [list(row) for row in workbook[title].iter_rows(max_col=7, values_only=True)]
            for title in fixture["sheets"]
        }
        if semantic_hash(sheets) != semantic_hash(cell_view(fixture)):
            raise ValueError("required_workbook_fixture_hash_mismatch")
        output[path.name] = semantic_hash(sheets)
    return output
